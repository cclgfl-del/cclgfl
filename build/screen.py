"""First-round screening of new submissions.

    NOTION_TOKEN=… NOTION_DATABASE_ID=… python build/screen.py

For every row with Status "Submitted" (or no Status yet) and an empty
Screening field:

1. Download the Word file and check it against the blog guidelines in site.yml:
   length (excluding notes), title and abstract length, number of authors,
   footnotes (not accepted), endnotes, and file name.
2. Look for the author's identity — in the text, in headers and footers, and in
   the file's hidden properties, comments and tracked changes.
3. Make an anonymised copy with the hidden identifying properties stripped, and
   attach it to the row. Reviewers open that copy, never the original.
4. Write a short report into Screening and move the row to "Screened".

Names found in the *text* are reported, not removed — rewriting an author's
prose is an editorial decision, not a script's.

Plagiarism and AI-content checks are not done here; there is no reliable free
tool for either, and a false positive in either direction does real harm.
"""

import os
import re
import sys

# Windows consoles default to a legacy code page; titles and arrows in the
# log would otherwise crash the run.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass

import tempfile
import zipfile
from pathlib import Path

import yaml

sys.path.insert(0, str(Path(__file__).parent))
import notion as N  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def words(s):
    return len(re.findall(r"[\w’'-]+", s or ""))


def names_from(authors):
    parts = re.split(r"\s*(?:,|;|\band\b|&)\s*", authors or "")
    return [p.strip() for p in parts if len(p.strip()) >= 4]


# ── reading the manuscript ────────────────────────────────────────────────

def inspect(path):
    """Everything screening needs from the .docx, read with python-docx plus a
    look at the raw parts python-docx does not expose (notes, comments)."""
    import docx

    d = docx.Document(str(path))
    body = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                body.append(cell.text)
    margins = []
    for s in d.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            try:
                margins.extend(p.text for p in part.paragraphs)
            except Exception:
                pass

    props = d.core_properties
    hidden = {"author": props.author or "", "last_modified_by": props.last_modified_by or "",
              "title": props.title or ""}

    footnotes = endnotes = 0
    people = set()
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        if "word/footnotes.xml" in names:
            footnotes = len(re.findall(r'<w:footnote [^>]*w:id="(?:[1-9]\d*)"', z.read("word/footnotes.xml").decode("utf8", "ignore")))
        if "word/endnotes.xml" in names:
            endnotes = len(re.findall(r'<w:endnote [^>]*w:id="(?:[1-9]\d*)"', z.read("word/endnotes.xml").decode("utf8", "ignore")))
        for part in ("word/document.xml", "word/comments.xml"):
            if part in names:
                people.update(re.findall(r'w:author="([^"]+)"', z.read(part).decode("utf8", "ignore")))
        if "docProps/app.xml" in names:
            company = re.search(r"<Company>([^<]*)</Company>", z.read("docProps/app.xml").decode("utf8", "ignore"))
            if company and company.group(1).strip():
                hidden["company"] = company.group(1).strip()

    return {"body": "\n".join(body), "margins": "\n".join(margins), "hidden": hidden,
            "footnotes": footnotes, "endnotes": endnotes, "people": sorted(people)}


def anonymise(src, dest):
    """Copy the .docx with identifying metadata removed: document properties,
    company, and the authors recorded on comments and tracked changes."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                s = data.decode("utf8")
                for tag in ("dc:creator", "cp:lastModifiedBy", "dc:title", "dc:subject", "cp:keywords", "dc:description"):
                    s = re.sub(r"<%s>.*?</%s>" % (tag, tag), "<%s></%s>" % (tag, tag), s, flags=re.S)
                data = s.encode("utf8")
            elif item.filename == "docProps/app.xml":
                s = data.decode("utf8")
                s = re.sub(r"<(Company|Manager)>.*?</\1>", r"<\1></\1>", s, flags=re.S)
                data = s.encode("utf8")
            elif item.filename in ("word/document.xml", "word/comments.xml", "word/footnotes.xml", "word/endnotes.xml"):
                s = data.decode("utf8")
                s = re.sub(r'w:author="[^"]*"', 'w:author="Reviewer"', s)
                s = re.sub(r'w:initials="[^"]*"', 'w:initials="R"', s)
                data = s.encode("utf8")
            zout.writestr(item, data)


# ── the report ────────────────────────────────────────────────────────────

def screen(info, page, filename, g):
    ok, flag = [], []
    title = N.text(page, "title")
    abstract = N.text(page, "abstract")
    authors = N.text(page, "authors")

    n = words(info["body"])
    (ok if g["words_min"] <= n <= g["words_max"] else flag).append(
        "%s words (guideline %s–%s, excluding notes)" % ("{:,}".format(n), "{:,}".format(g["words_min"]), "{:,}".format(g["words_max"])))

    tw = words(title)
    (ok if tw <= g["title_max_words"] else flag).append("Title %d words (max %d)" % (tw, g["title_max_words"]))

    if abstract:
        aw = words(abstract)
        (ok if aw <= g["abstract_max_words"] else flag).append("Abstract %d words (max %d)" % (aw, g["abstract_max_words"]))
    else:
        flag.append("No abstract supplied")

    count = len(names_from(authors)) or (1 if authors else 0)
    (ok if 0 < count <= g["max_authors"] else flag).append("%d author(s) (max %d)" % (count, g["max_authors"]))

    if info["footnotes"] and not g.get("footnotes_allowed", False):
        flag.append("%d footnote(s) — footnotes are not accepted; hyperlinks or endnotes only" % info["footnotes"])
    if info["endnotes"]:
        ok.append("%d endnote(s)" % info["endnotes"])

    if not re.search(r"_CCLGFL Blog\.docx$", filename, re.I):
        flag.append("File name should follow “%s”" % g["filename_pattern"])

    # Identity
    names = names_from(authors)
    found_text = [nm for nm in names if nm.casefold() in info["body"].casefold()]
    found_margins = [nm for nm in names if nm.casefold() in info["margins"].casefold()]
    hidden = [("%s: %s" % (k, v)) for k, v in info["hidden"].items() if v and any(
        part.casefold() in v.casefold() for nm in names for part in nm.split() if len(part) >= 3)]
    if found_text:
        flag.append("Author name in the text: %s — not removed, please check" % ", ".join(found_text))
    if found_margins:
        flag.append("Author name in a header or footer: %s — not removed, please check" % ", ".join(found_margins))
    if hidden or info["people"]:
        ok.append("Identifying file properties removed in the anonymised copy")

    lines = ["Screened automatically against the blog guidelines."]
    lines += ["⚠ " + f for f in flag]
    lines += ["✓ " + o for o in ok]
    return "\n".join(lines), flag


def main():
    if not os.environ.get("NOTION_TOKEN") or not os.environ.get("NOTION_DATABASE_ID"):
        print("Screening skipped: NOTION_TOKEN / NOTION_DATABASE_ID not set.")
        return
    g = yaml.safe_load((ROOT / "site.yml").read_text(encoding="utf-8"))["guidelines"]
    client = N.Notion()
    # A form response arrives with Status empty; treat that the same as
    # "Submitted", so the form needs no special configuration.
    rows = client.query(os.environ["NOTION_DATABASE_ID"], filter={"and": [
        {"or": [
            {"property": N.FIELDS["status"], "select": {"equals": "Submitted"}},
            {"property": N.FIELDS["status"], "select": {"is_empty": True}},
        ]},
        {"property": N.FIELDS["screening"], "rich_text": {"is_empty": True}},
    ]})
    print("Screening %d new submission(s)" % len(rows))

    for page in rows:
        title = N.text(page, "title") or page["id"]
        docs = [f for f in N.files(page, "manuscript") if f[0].lower().endswith((".docx", ".doc"))]
        if not docs:
            client.update(page["id"], {N.FIELDS["screening"]: N.rich_text_value("⚠ No Word file attached.")})
            print("  %s: no manuscript" % title)
            continue
        name, url = docs[-1]
        if name.lower().endswith(".doc"):
            client.update(page["id"], {N.FIELDS["screening"]: N.rich_text_value(
                "⚠ Legacy .doc file — automatic checks need .docx. Ask the author to resave, or screen by hand.")})
            print("  %s: legacy .doc" % title)
            continue

        with tempfile.TemporaryDirectory() as tmp:
            local = client.download(url, Path(tmp) / "manuscript.docx")
            report, flags = screen(inspect(local), page, name, g)

            anon = Path(tmp) / ("%s (anonymised).docx" % re.sub(r"[^\w\s-]", "", title)[:60].strip())
            anonymise(local, anon)
            props = {N.FIELDS["screening"]: N.rich_text_value(report),
                     N.FIELDS["status"]: {"select": {"name": "Screened"}}}
            try:
                upload_id = client.upload_file(anon, DOCX)
                props[N.FIELDS["anonymised"]] = {"files": [
                    {"type": "file_upload", "file_upload": {"id": upload_id}, "name": anon.name}]}
            except N.NotionError as e:
                props[N.FIELDS["screening"]] = N.rich_text_value(
                    report + "\n⚠ Could not attach the anonymised copy (%s) — strip file properties by hand." % e)
            client.update(page["id"], props)
        print("  %s: %d flag(s)" % (title, len(flags)))


if __name__ == "__main__":
    main()
