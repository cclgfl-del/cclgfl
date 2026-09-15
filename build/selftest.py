"""Offline checks for the pipeline — no Notion token needed.

    python build/selftest.py

Builds a real Word manuscript with pandoc, then runs it through conversion,
screening and anonymisation; converts a set of Notion-shaped blocks; and
renders both to confirm notes, paragraph numbers and heading numerals.
Exits non-zero on the first failure.
"""

import sys

# Windows consoles default to a legacy code page; titles and arrows in the
# log would otherwise crash the run.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass

import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import content  # noqa: E402
import manuscript  # noqa: E402
import screen  # noqa: E402

PASS = []


def check(name, cond, detail=""):
    if not cond:
        print("FAIL  %s %s" % (name, detail))
        sys.exit(1)
    PASS.append(name)
    print("ok    %s" % name)


def fake_page(**props):
    def rt(s):
        return {"type": "rich_text", "rich_text": [{"plain_text": s}]}
    return {"id": "test", "properties": {
        "Title": {"type": "title", "title": [{"plain_text": props.get("title", "")}]},
        "Abstract": rt(props.get("abstract", "")),
        "Authors": rt(props.get("authors", "")),
    }}


def main():
    import docx
    import pypandoc

    tmp = Path(tempfile.mkdtemp())

    # ── a manuscript the way authors actually write them ─────────────────
    source = (
        "**Why the threshold matters**\n\n"
        "**I. Introduction**\n\n"
        "The first paragraph cites a case.[^1] Priya Sharma wrote this sentence.\n\n"
        "**A. The narrower point**\n\n"
        "A second paragraph with [a link](https://www.sebi.gov.in/).\n\n"
        + ("Filler sentence for the word count check. " * 30) + "\n\n"
        "[^1]: *Swiss Ribbons Pvt. Ltd. v. Union of India*, (2019) 4 SCC 17.\n"
    )
    md_in = tmp / "in.md"
    md_in.write_text(source, encoding="utf-8")
    docx_path = tmp / "Why the threshold matters_CCLGFL Blog.docx"
    pypandoc.convert_file(str(md_in), "docx", outputfile=str(docx_path))
    d = docx.Document(str(docx_path))
    d.core_properties.author = "Priya Sharma"
    d.core_properties.last_modified_by = "Priya Sharma"
    d.save(str(docx_path))

    # conversion
    post_dir = tmp / "post"
    md = manuscript.docx_to_markdown(docx_path, post_dir, title="Why the threshold matters")
    check("title line removed from body", "Why the threshold matters" not in md.split("\n")[0], md[:120])
    check("bold roman line becomes an h2", "## I. Introduction" in md, md[:300])
    check("bold lettered line becomes an h3", "### A. The narrower point" in md, md[:400])
    check("note survives conversion", "[^1]:" in md and "Swiss Ribbons" in md)

    r = content.render(md)
    check("note rendered as a sidenote", r["notes"] == 1 and 'class="sn"' in r["html"])
    check("heading numeral split out", '<span class="h-no">I</span>' in r["html"])
    check("paragraphs numbered", r["paragraphs"] >= 3 and 'id="p1"' in r["html"])
    check("sidenote sits inside its paragraph", r["html"].index("sn-1") < r["html"].index("</p>", r["html"].index("sn-1")))

    # screening
    info = screen.inspect(docx_path)
    check("footnote counted", info["footnotes"] == 1, str(info["footnotes"]))
    check("hidden author property found", info["hidden"]["author"] == "Priya Sharma")
    page = fake_page(title="Why the threshold matters", abstract="Short abstract.", authors="Priya Sharma")
    report, flags = screen.screen(info, page, docx_path.name, {
        "words_min": 1000, "words_max": 1500, "title_max_words": 10,
        "max_authors": 2, "hyperlinks_only": True, "filename_pattern": "Title_CCLGFL Blog"})
    check("short manuscript flagged", any("words" in f for f in flags), report)
    check("footnotes flagged", any("footnote" in f for f in flags), report)
    check("author name in text flagged", any("Author name in the text" in f for f in flags), report)
    check("file name accepted", not any("File name" in f for f in flags), report)

    # anonymisation
    anon = tmp / "anon.docx"
    screen.anonymise(docx_path, anon)
    with zipfile.ZipFile(anon) as z:
        core = z.read("docProps/core.xml").decode("utf8")
    check("creator stripped from anonymised copy", "Priya Sharma" not in core)
    check("anonymised copy still opens", len(docx.Document(str(anon)).paragraphs) > 0)

    # ── a post written directly in Notion ────────────────────────────────
    def para(text, **ann):
        return {"type": "paragraph", "has_children": False,
                "paragraph": {"rich_text": [{"plain_text": text, "annotations": ann}]}}

    def heading(text, level=1):
        k = "heading_%d" % level
        return {"type": k, "has_children": False, k: {"rich_text": [{"plain_text": text}]}}

    def item(kind, text):
        return {"type": kind, "has_children": False, kind: {"rich_text": [{"plain_text": text}]}}

    blocks = [
        para("Opening paragraph with a note.[^1] And a second.[^2]"),
        heading("II. Analysis"),
        item("bulleted_list_item", "first point"),
        item("bulleted_list_item", "second point"),
        para("Paragraph after a list."),
        heading("Notes"),
        item("numbered_list_item", "The first note."),
        item("numbered_list_item", "The second note."),
    ]
    md2 = manuscript.blocks_to_markdown(blocks, lambda _id: [], lambda url, n: "x.png")
    check("notes list becomes definitions", "[^1]: The first note." in md2 and "[^2]: The second note." in md2, md2)
    check("note markers not escaped", "note.[^1]" in md2, md2)
    check("list closed before next paragraph", "- second point\n\nParagraph after a list." in md2, md2)
    r2 = content.render(md2)
    check("Notion notes render as sidenotes", r2["notes"] == 2, str(r2["notes"]))
    check("'Notes' heading not rendered", ">Notes<" not in r2["html"])

    sync_checks(tmp)
    print("\n%d checks passed" % len(PASS))


class FakeNotion:
    """Stands in for the Notion API: serves pages and blocks from memory and
    records every write, so the sync's decisions can be checked exactly."""

    def __init__(self, pages, blocks):
        self.pages, self.blocks = pages, blocks
        self.updates, self.downloads = [], 0

    def query(self, db_id, filter=None, sorts=None):
        wanted = {c["select"]["equals"] for c in filter["or"]}
        return [p for p in self.pages if p["properties"]["Status"]["select"]["name"] in wanted]

    def children(self, block_id):
        return self.blocks.get(block_id, [])

    def update(self, page_id, properties):
        self.updates.append((page_id, properties))

    def download(self, url, dest):
        self.downloads += 1
        Path(dest).parent.mkdir(parents=True, exist_ok=True)
        Path(dest).write_bytes(b"x")
        return Path(dest)


def notion_page(pid, title, status, edited="2026-09-01T10:00:00.000Z", slug=""):
    rt = lambda s: {"type": "rich_text", "rich_text": [{"plain_text": s}] if s else []}  # noqa: E731
    return {"id": pid, "last_edited_time": edited, "properties": {
        "Title": {"type": "title", "title": [{"plain_text": title}]},
        "Status": {"type": "select", "select": {"name": status}},
        "Authors": rt("A. Author"), "Author bio": rt(""), "Abstract": rt("An abstract."),
        "Slug": rt(slug), "Manuscript": {"type": "files", "files": []}, "Header image": {"type": "files", "files": []},
        "Publish date": {"type": "date", "date": {"start": "2026-09-01"}},
        "Featured": {"type": "checkbox", "checkbox": False},
        "Screening": rt(""), "Preview link": {"type": "url", "url": None}, "Live link": {"type": "url", "url": None},
    }}


def sync_checks(tmp):
    import yaml
    import notion_sync as S

    root = tmp / "site"
    S.POSTS, S.PREVIEWS = root / "posts", root / "previews"
    cfg = yaml.safe_load((Path(__file__).resolve().parent.parent / "site.yml").read_text(encoding="utf-8"))
    cfg["site_url"] = "https://example.test"

    body = [{"type": "paragraph", "has_children": False,
             "paragraph": {"rich_text": [{"plain_text": "Body text."}]}}]
    pub = notion_page("p-live", "A published post", "Published")
    acc = notion_page("p-draft", "An accepted draft", "Accepted")
    api = FakeNotion([pub, acc], {"p-live": body, "p-draft": body})

    changed, errors = S.sync(api, "db", cfg)
    check("sync: first run reports a change", changed and not errors, str(errors))
    check("sync: published post written to posts/", (S.POSTS / "a-published-post" / "index.md").exists())
    token = S.preview_token("p-draft")
    check("sync: accepted post written to previews/, not posts/",
          (S.PREVIEWS / token / "index.md").exists() and not (S.POSTS / "an-accepted-draft").exists())
    check("sync: preview fingerprint committed", (S.POSTS.parent / ".previews-state").exists())
    written = {(pid, list(props)[0]) for pid, props in api.updates}
    check("sync: live and preview links written back", ("p-live", "Live link") in written and ("p-draft", "Preview link") in written, str(written))

    api.updates.clear()
    changed, _ = S.sync(api, "db", cfg)
    check("sync: unchanged Notion means no change", not changed)

    pub["last_edited_time"] = "2026-09-02T10:00:00.000Z"
    pub["properties"]["Slug"] = {"type": "rich_text", "rich_text": [{"plain_text": "renamed-post"}]}
    changed, _ = S.sync(api, "db", cfg)
    check("sync: edited slug moves the post", changed and (S.POSTS / "renamed-post").exists()
          and not (S.POSTS / "a-published-post").exists())

    bad = notion_page("p-bad", "", "Published")
    api.pages.append(bad)
    api.blocks["p-bad"] = body
    changed, errors = S.sync(api, "db", cfg)
    check("sync: a bad row is reported, not fatal", errors and (S.POSTS / "renamed-post").exists())
    check("sync: the reason is written onto the Notion page",
          any(pid == "p-bad" and "Screening" in props for pid, props in api.updates))

    api.pages = [acc]
    try:
        S.sync(api, "db", cfg)
        refused = False
    except SystemExit:
        refused = True
    check("sync: refuses to unpublish every post at once", refused)


if __name__ == "__main__":
    main()
